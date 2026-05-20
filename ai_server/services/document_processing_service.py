#
#  @FileName : document_processing_service.py
#  @Description : 문서 본문 추출/청크 분할/벡터 생성 서비스 모듈
#  @Author : 김다솜
#  @Date : 2026. 05. 12
#  @Modification_History
#  @
#  @ 수정일자        수정자          수정내용
#  @ ----------    ---------    -------------------------------
#  @ 2026.05.12    김다솜        문서 자동 처리 파이프라인 구현, 원격 문서 요청 timeout/User-Agent 보강 및 상단 주석 보강
#  @ 2026.05.20    송혜진        Google Drive 파일 ID 추출/폴더 링크 제외/DOCX 추출 보강.
#

import hashlib
import html
import json
import math
import re
import zipfile
from io import BytesIO
from pathlib import Path
from urllib.parse import parse_qs, urlparse

import pdfplumber
import requests
from pypdf import PdfReader
from docx import Document as DocxDocument

from services.ollama_client import summarize_document
from schemas.document_schema import (
    DocumentChunkResponse,
    DocumentProcessRequest,
    DocumentProcessResponse,
)

# --- 글로벌 설정 상수 ---
MAX_CHUNK_CHARS = 900  # 하나의 청크에 들어갈 최대 글자 수
EMBED_DIMENSION = 256  # 생성할 해시 임베딩 벡터의 차원 수
REMOTE_CONNECT_TIMEOUT = 10 # 원격 서버 연결 제한 시간(초)
REMOTE_READ_TIMEOUT = 90    # 원격 서버 데이터 읽기 제한 시간(초)

# 웹 크롤링/ 다운로드 시 차단을 방지하기 위한 표준 브라우저 User-Agent 및 헤더 설정
REMOTE_REQUEST_HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/124.0 Safari/537.36"
    ),
    "Accept": "text/html,application/pdf,text/plain,application/json,*/*",
}

# 구글 드라이브 및 구글 독스 URL 분석을 위한 정규식 패턴
GOOGLE_DRIVE_FOLDER_PATTERNS = (
    r"/drive/folders/",
    r"/drive/u/0/folders/",
)
GOOGLE_DRIVE_FILE_PATTERN = re.compile(r"/file/d/([^/]+)/")
GOOGLE_DRIVE_OPEN_PATTERN = re.compile(r"[?&]id=([^&]+)")
GOOGLE_DOCS_FILE_PATTERN = re.compile(r"/document/d/([^/]+)/")


def process_document(req: DocumentProcessRequest) -> DocumentProcessResponse:
    """
    [RAG 핵심 파이프라인 메인 함수]
    문서 경로(로컬/원격)를 받아 텍스트 추출 -> 정제 -> 청크 분할 -> 벡터 생성
    """
    # [1] 문서 위치(로컬/웹) 및 확장자에 따른 본문 텍스트 추출
    text, source_type = extract_text(req.filePath)

    # [2] 텍스트 정제 (불필요한 공백 제거, 제어 문자 제거 등)
    cleaned_text = normalize_text(text)

    if not cleaned_text:
        raise ValueError("문서에서 추출한 본문이 비어 있습니다.")

    # [3] PDF의 경우, 텍스트 가독성/ 품질 검사 진행 (텍스트가 깨졌거나 이미지만 있는 경우 방지)
    if source_type.endswith("pdf") and is_low_quality_extraction(cleaned_text):
        raise ValueError("PDF 본문 추출 품질이 너무 낮습니다. 다른 PDF를 사용하거나 OCR 처리가 필요합니다.")

    # [4] 정제 된 본문을 문맥 단락 기준으로 분할(청킹)
    chunks = split_text(cleaned_text)
    if not chunks:
        raise ValueError("청크를 생성할 수 없습니다.")

    # [5] Ollama 요약 모델을 통해 문서 프리뷰 생성 (실패 시 앞 300자 커트)
    preview_text = build_preview_text(cleaned_text)

    # [6] 각 청크별 메타데이터 구성 및 알고리즘 기반 해시 백터(Embedding) 생성
    response_chunks: list[DocumentChunkResponse] = []
    for index, chunk_text in enumerate(chunks, start=1):
        token_count = estimate_token_count(chunk_text)               # 예측 토큰 수 계산
        embedding = make_hash_embedding(chunk_text, EMBED_DIMENSION) # 경량 알고리즘 임베딩 생성

        response_chunks.append(
            DocumentChunkResponse(
                chunkNo=index,
                content=chunk_text,
                tokenCount=token_count,
                sectionTitle=build_section_title(req.title, index, chunk_text), # 소제목 자동 생성
                embeddingData=json.dumps(embedding),
                modelName="hash-embedding-v1",
                dimension=EMBED_DIMENSION,
            )
        )

    # [7] 최종 가공된 최종 스키마 리턴
    return DocumentProcessResponse(
        sourceType=source_type,
        extractedTextPreview=preview_text[:300],
        chunkCount=len(response_chunks),
        vectorCount=len(response_chunks),
        chunks=response_chunks,
    )


def extract_text(file_path: str) -> tuple[str, str]:
    """
    경로 형식을 판단하여 로컬 파일 시스템 혹은 원격 웹 URL에서 텍스트를 추출
    """
    # HTTP/ HTTPS로 시작하면 웹 브라우징/ 원격 다운로드 다운로드 프로세스 수행
    if file_path.startswith(("http://", "https://")):
        return extract_remote_text(file_path)

    # 로컬 파일 확장자별 수기 처리
    lower_path = file_path.lower()
    if lower_path.endswith(".pdf"):
        return extract_local_pdf_text(file_path), "local-pdf"

    if lower_path.endswith(".docx"):
        return extract_local_docx_text(file_path), "local-docx"

    if lower_path.endswith((".html", ".htm")):
        with open(file_path, "r", encoding="utf-8") as file:
            return html_to_text(file.read()), "local-html"

    # 기본 텍스트 파일 (.txt, .md 등) 처리
    with open(file_path, "r", encoding="utf-8") as file:
        return file.read(), "local-text"


def extract_remote_text(file_path: str) -> tuple[str, str]:
    """
    원격 URL 주소를 분석하여 일반 웹페이지, 구글 드라이브, 구글 독스 여부를 가려 텍스트 가져오기.
    """
    # 폴더 통째 공유 링크는 지원 배제
    if is_google_drive_folder_url(file_path):
        raise ValueError("Google Drive 폴더 링크는 현재 지원하지 않습니다. 공개 파일 공유 링크를 입력해 주세요.")

    # 구글 독스(온라인 문서 작성기) 문서 처리
    if is_google_docs_url(file_path):
        file_id = extract_google_drive_file_id(file_path)
        if not file_id:
            raise ValueError("Google Docs 문서 ID를 추출할 수 없습니다.")
        return extract_google_docs_text(file_id)

    # 구글 드라이브 단일 파일 공유 링크 처리
    if is_google_drive_url(file_path):
        file_id = extract_google_drive_file_id(file_path)
        if not file_id:
            raise ValueError("Google Drive 파일 ID를 추출할 수 없습니다.")
        return extract_remote_drive_file(file_path, file_id)

    # 그 외 일반 다운로드 링크 및 크롤링 대상 처리
    return extract_generic_remote_text(file_path)


def extract_local_pdf_text(file_path: str) -> str:
    """로컬 PDF 파일을 바이너리로 읽어 텍스트 추출부로 전달."""
    with open(file_path, "rb") as file:
        return extract_pdf_text(file.read())


def extract_local_docx_text(file_path: str) -> str:
    """로컬 Word(.docx) 파일을 바이너리로 읽어 텍스트 추출부로 전달."""
    with open(file_path, "rb") as file:
        return extract_docx_text(file.read())


def extract_pdf_text(pdf_bytes: bytes) -> str:
    """
    [이중 추출 구조]
    속도가 빠른 pypdf로 1차 추출을 시도하고, 
    결과 품질이 낮을 경우 정밀한 pdfplumber로 2차 폴백(Fallback) 추출을 진행.
    """
    # 1차 추출 시도
    text = extract_pdf_text_with_pypdf(pdf_bytes)

    # 1차 결과가 나쁠 시(글자 깨짐, 스페이스 오류 등) 2차 엔진 가동
    if is_low_quality_extraction(text):
        fallback_text = extract_pdf_text_with_pdfplumber(pdf_bytes)
        if not is_low_quality_extraction(fallback_text):
            text = fallback_text

    if not text:
        raise ValueError("PDF에서 추출한 본문이 비어 있습니다.")

    return text


def extract_pdf_text_with_pypdf(pdf_bytes: bytes) -> str:
    """pypdf 라이브러리를 활용한 고속 PDF 텍스트 파싱"""
    reader = PdfReader(BytesIO(pdf_bytes))
    pages: list[str] = []

    for page in reader.pages:
        page_text = page.extract_text() or ""
        page_text = page_text.strip()
        if page_text:
            pages.append(page_text)

    return "\n\n".join(pages).strip()


def extract_pdf_text_with_pdfplumber(pdf_bytes: bytes) -> str:
    """pdfplumber 라이브러리를 활용한 정밀 PDF 레이아웃 분석 및 텍스트 파싱"""
    pages: list[str] = []

    with pdfplumber.open(BytesIO(pdf_bytes)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            page_text = page_text.strip()
            if page_text:
                pages.append(page_text)

    return "\n\n".join(pages).strip()


def extract_docx_text(docx_bytes: bytes) -> str:
    """
    python-docx를 이용해
    Word 문서 내 일반 문단(Paragraph)과 
    표(Table) 안의 데이터 텍스트를 구조적으로 결합해 추출합니다.
    """
    try:
        document = DocxDocument(BytesIO(docx_bytes))
    except Exception as exc:
        raise ValueError("DOCX 본문 추출에 실패했습니다.") from exc

    parts: list[str] = []

    # [1] 일반 문단 텍스트 수집
    for paragraph in document.paragraphs:
        paragraph_text = (paragraph.text or "").strip()
        if paragraph_text:
            parts.append(paragraph_text)

    # [2] 표 내부 텍스트 수집 (열 단위를 파이프 기회 '|'로 연결해 데이터 맥락 보존)
    for table in document.tables:
        for row in table.rows:
            cell_texts = [
                (cell.text or "").strip()
                for cell in row.cells
                if (cell.text or "").strip()
            ]
            if cell_texts:
                parts.append(" | ".join(cell_texts))

    text = "\n\n".join(parts).strip()
    if not text:
        raise ValueError("DOCX에서 추출 가능한 텍스트가 없습니다.")

    return text


def extract_generic_remote_text(request_url: str) -> tuple[str, str]:
    """일반 원격 웹 페이지 혹은 다이렉트 파일 다운로드 주소로부터 데이터 요청"""
    response = fetch_remote_response(request_url, request_url)
    return extract_response_payload(response, request_url, "remote")


def extract_remote_drive_file(source_url: str, file_id: str) -> tuple[str, str]:
    """구글 드라이브 단일 파일의 고유 ID를 다운로드 엔드포인트 URL로 조립하여 다이렉트 패치"""
    download_url = build_google_drive_download_url(file_id)
    response = fetch_remote_response(download_url, source_url)
    return extract_response_payload(response, source_url, "remote-drive")


def extract_google_docs_text(file_id: str) -> tuple[str, str]:
    """
    구글 독스는 일반 다운로드가 안 되므로 구글 수출(Export) API 사용.
    1차로 가벼운 TXT로 내보내기를 시도하고 권한/형식 문제 발생 시 2차로 HTML 포맷 내보내기를 시도.
    """
    # 1차 시도 : TXT 포맷 다운로드
    txt_export_url = f"https://docs.google.com/document/d/{file_id}/export?format=txt"
    response = fetch_remote_response(txt_export_url, txt_export_url)
    text = decode_response_text(response).strip()
    if text and not is_google_drive_access_denied_text(text):
        return text, "remote-google-docs-txt"

    # 2차 폴백 시도 : HTML 포맷 다운로드 후 태그 제거
    html_export_url = f"https://docs.google.com/document/d/{file_id}/export?format=html"
    response = fetch_remote_response(html_export_url, html_export_url)
    html_text = html_to_text(decode_response_text(response))
    if html_text.strip():
        return html_text, "remote-google-docs-html"

    raise ValueError("Google Docs 본문을 추출할 수 없습니다.")


def extract_response_payload(response: requests.Response, source_url: str, source_prefix: str) -> tuple[str, str]:
    """
    HTTP 응답 객체의 Content-Type 헤더와 파일 시그니처바이트(Magic Byte)를 검사하여 
    동적으로 맞는 확장자별 추출 엔진(PDF, DOCX, HTML, TEXT)을 호출합니다.
    """
    content_type = (response.headers.get("content-type") or "").lower()
    filename = extract_filename_from_response(response, source_url)
    detected_kind = detect_remote_file_kind(response.content, content_type, filename, response.url or source_url)

    if detected_kind == "pdf":
        return extract_pdf_text(response.content), f"{source_prefix}-pdf"

    if detected_kind == "docx":
        return extract_docx_text(response.content), f"{source_prefix}-docx"

    if detected_kind == "html":
        return html_to_text(decode_response_text(response)), f"{source_prefix}-html"

    if detected_kind == "text":
        return decode_response_text(response), f"{source_prefix}-text"

    # 알 수 없는 형식일 때 최종적으로 HTML 구조나 일반 텍스트 포맷인지 문자열을 유추해 파싱
    decoded = decode_response_text(response)
    if looks_like_html(decoded):
        return html_to_text(decoded), f"{source_prefix}-html"

    if decoded.strip():
        return decoded, f"{source_prefix}-text"

    raise ValueError("지원하지 않는 파일 형식입니다.")


def fetch_remote_response(request_url: str, source_url: str) -> requests.Response:
    """
    requests 라이브러리로 원격지 데이터를 패치하며 타임아웃 예외, 
    구글 드라이브 비공개 권한 거부, 대용량 파일 경고 페이지 유무 등을 통합 검증합니다.
    """
    try:
        response = requests.get(
            request_url,
            headers=REMOTE_REQUEST_HEADERS,
            timeout=(REMOTE_CONNECT_TIMEOUT, REMOTE_READ_TIMEOUT),
        )
    except requests.Timeout as exc:
        raise TimeoutError(
            f"원격 문서 응답 시간이 {REMOTE_READ_TIMEOUT}초를 초과했습니다. "
            "공개 파일 공유 링크인지 확인하거나 파일을 로컬/드라이브 문서로 등록해 주세요."
        ) from exc
    except requests.RequestException as exc:
        raise ValueError(f"파일 다운로드에 실패했습니다: {exc}") from exc

    # 401, 403 권한 에러 방어
    if response.status_code in (401, 403):
        raise ValueError("비공개 또는 권한이 필요한 Google Drive 링크입니다.")

    response.raise_for_status()

    # 구글 로그인 페이지로 리다이렉트 되는 경우 권한 오류로 판정
    if is_google_auth_redirect(response.url):
        raise ValueError("비공개 또는 권한이 필요한 Google Drive 링크입니다.")

    content_type = (response.headers.get("content-type") or "").lower()
    decoded_text = None
    if content_type.startswith("text/") or "html" in content_type:
        decoded_text = decode_response_text(response)

    # 대용량 바이러스 검사 불가 경고창 토큰 체크 방어
    if decoded_text and is_google_drive_confirmation_page(decoded_text):
        raise ValueError("Google Drive 대용량 파일의 확인 토큰이 필요한 링크는 지원하지 않습니다.")

    # 텍스트 내용 중 접근 거부 메시지가 포함되어 있는지 검사
    if decoded_text and is_google_drive_access_denied_text(decoded_text):
        raise ValueError("비공개 또는 권한이 필요한 Google Drive 링크입니다.")

    return response


def build_google_drive_download_url(file_id: str) -> str:
    """구글 드라이브 단일 파일 다이렉트 다운로드 주소를 완성."""
    return f"https://drive.google.com/uc?export=download&id={file_id}"


def is_google_drive_url(file_path: str) -> bool:
    """호스트 도메인이 구글 드라이브 주소인지 검사."""
    parsed = urlparse(file_path)
    return "drive.google.com" in parsed.netloc


def is_google_docs_url(file_path: str) -> bool:
    """호스트 도메인이 구글 독스 편집기 주소인지 검사."""
    parsed = urlparse(file_path)
    return "docs.google.com" in parsed.netloc and "/document/d/" in parsed.path


def is_google_drive_folder_url(file_path: str) -> bool:
    """경로명 분석을 통해 파일이 아닌 대형 폴더 공유 주소인지 판별."""
    parsed = urlparse(file_path)
    if "drive.google.com" not in parsed.netloc:
        return False

    return any(pattern in parsed.path for pattern in GOOGLE_DRIVE_FOLDER_PATTERNS)


def extract_google_drive_file_id(file_path: str) -> str | None:
    """다양한 구글 공유 URL 서식 파라미터나 REST Path 세그먼트 사이에서 고유 파일 ID를 추출."""
    parsed = urlparse(file_path)
    query = parse_qs(parsed.query)

    # [1] 쿼리스트링 파라미터(?id=xxx) 방식 검사
    if "id" in query and query["id"]:
        return query["id"][0]

    # [2] URL 경로 (/file/d/xxx/) 패턴 검사
    if match := GOOGLE_DRIVE_FILE_PATTERN.search(file_path):
        return match.group(1)

    # [3] 구글 독스 문서 주소 (/document/d/xxx/) 패턴 검사
    if match := GOOGLE_DOCS_FILE_PATTERN.search(file_path):
        return match.group(1)

    # [4] 오픈 주소 매칭 검사
    if match := GOOGLE_DRIVE_OPEN_PATTERN.search(file_path):
        return match.group(1)

    return None


def extract_filename_from_response(response: requests.Response, request_url: str) -> str:
    """HTTP 헤더의 Content-Disposition 파트에서 오리지널 파일명을 파싱하고 없으면 URL 끝자리 파일명을 가져오기."""
    content_disposition = response.headers.get("content-disposition", "")
    if content_disposition:
        match = re.search(
            r'filename\*=(?:UTF-8\'\')?"?([^\";]+)"?|filename="?([^\";]+)"?',
            content_disposition,
            re.IGNORECASE,
        )
        if match:
            filename = match.group(1) or match.group(2)
            if filename:
                return filename.strip().strip('"')

    parsed_url = urlparse(request_url)
    return Path(parsed_url.path).name or ""


def detect_remote_file_kind(
    content: bytes,
    content_type: str,
    filename: str,
    request_url: str,
) -> str:
    """바이너리 데이터의 실측 구조정보와 헤더, 확장자 힌트를 총동원해 문서 타입을 정확히 매핑."""
    filename_lower = (filename or "").lower()
    url_lower = (request_url or "").lower()

    # pdf 조건 분기
    if is_pdf_bytes(content) or filename_lower.endswith(".pdf") or "application/pdf" in content_type:
        return "pdf"

    # word 파일 조건 분기
    if is_docx_bytes(content) or filename_lower.endswith(".docx") or "wordprocessingml.document" in content_type:
        return "docx"

    # html 구조 분기
    if "text/html" in content_type or filename_lower.endswith((".html", ".htm")) or url_lower.endswith((".html", ".htm")):
        return "html"

    # 일반 텍스트 포맷 분기
    if (
        "text/plain" in content_type
        or "application/json" in content_type
        or "text/csv" in content_type
        or filename_lower.endswith((".txt", ".md", ".csv", ".json"))
    ):
        return "text"

    # 헤더나 확장자가 없거나 모호할 때 바이너리 4000자 초반 내용을 읽어 문ㅇ자열 패턴 분석 유추
    preview = decode_bytes_preview(content)
    if looks_like_html(preview):
        return "html"

    if looks_like_text(preview):
        return "text"

    return "unknown"


def is_pdf_bytes(content: bytes) -> bool:
    """바이너리 매직 넘버가 PDF 표준 규격 코드인 %PDF- 로 시작하는지 검증."""
    return bool(content and content.startswith(b"%PDF-"))


def is_docx_bytes(content: bytes) -> bool:
    """DOCX 압축 파일 바이너리를 열어 필수 규격 구조 XML 파일들이 온전히 들어있는지 확인"""
    if not content:
        return False

    try:
        if not zipfile.is_zipfile(BytesIO(content)):
            return False

        with zipfile.ZipFile(BytesIO(content)) as zip_file:
            names = set(zip_file.namelist())
            # DOCX 포맷을 증명하는 필수 아키텍쳐 파일 엔트리 파일 체크
            return "[Content_Types].xml" in names and "word/document.xml" in names
    except zipfile.BadZipFile:
        return False


def decode_bytes_preview(content: bytes, max_length: int = 4000) -> str:
    """바이너리 전반부 버퍼를 복수의 인코딩 스키마를 순회 대입하며 안전한 문자열 프리뷰로 변환."""
    if not content:
        return ""

    preview = content[:max_length]
    for encoding in ("utf-8", "cp949", "latin1"):
        try:
            return repair_mojibake_text(preview.decode(encoding, errors="ignore"))
        except LookupError:
            continue

    return repair_mojibake_text(preview.decode("utf-8", errors="ignore"))


def looks_like_html(text: str) -> bool:
    """텍스트의 구조가 HTML 도큐먼트 형태를 띠고 있는지 여부 검사."""
    candidate = (text or "").strip().lower()
    return candidate.startswith("<!doctype html") or candidate.startswith("<html") or "<body" in candidate


def looks_like_text(text: str) -> bool:
    """프린트 가능 문자 공간 분포율을 계산하여 바이너리가 이진 데이터가 아닌 문자 데이터인지 분석"""
    candidate = (text or "").strip()
    if not candidate:
        return False

    printable = sum(1 for ch in candidate if ch.isprintable() or ch.isspace())
    return printable / max(1, len(candidate)) > 0.8


def is_google_drive_confirmation_page(text: str) -> bool:
    """구글 드라이브 대용량 다운로드 시 나타나는 '바이러스 검사 불가 경고 페이지' 유무를 판별."""
    candidate = (text or "").lower()
    return "confirm=" in candidate or "download_warning" in candidate or "too large to scan" in candidate


def is_google_drive_access_denied_text(text: str) -> bool:
    """HTML 본문 내에 구글 권한 제한 오류 토큰어구가 존재하는지 검증."""
    candidate = (text or "").lower()
    tokens = [
        "access denied",
        "permission",
        "need access",
        "sign in",
        "로그인",
        "권한",
    ]
    return any(token in candidate for token in tokens)


def is_google_auth_redirect(url: str) -> bool:
    """구글 계정 로그인 권한 챌린지 페이지로 화면 주소가 튕겼는지 체크"""
    candidate = (url or "").lower()
    return "accounts.google.com" in candidate or "servicelogin" in candidate or "signin" in candidate


def html_to_text(raw_html: str) -> str:
    """정규식을 활용한 HTML 페이지 내부의 스크립트, CSS 스타일 태그 블록을 소거하고 알맹이 텍스트만 남김"""
    without_script = re.sub(r"<script.*?>.*?</script>", " ", raw_html, flags=re.IGNORECASE | re.DOTALL)
    without_style = re.sub(r"<style.*?>.*?</style>", " ", without_script, flags=re.IGNORECASE | re.DOTALL)
    without_tags = re.sub(r"<[^>]+>", " ", without_style)
    return html.unescape(without_tags) # 이스케이프 엔티티 복원


def decode_response_text(response: requests.Response) -> str:
    """HTTP 응답 인코딩 선언에 맞춰 텍스트를 디코딩하고 예외 시 utf-8 강제 디코딩을 실행."""
    encoding = response.encoding or response.apparent_encoding or "utf-8"

    try:
        text = response.content.decode(encoding, errors="ignore")
    except LookupError:
        text = response.content.decode("utf-8", errors="ignore")

    return repair_mojibake_text(text)


def normalize_text(text: str) -> str:
    """RAG 임베딩 성능 정하를 방지하기 위해 널 문자 제거, 개행 표준화, 연속 공백 일축 등 텍스트 표준화를 수행"""
    text = repair_mojibake_text(text)
    text = text.replace("\x00", " ")
    text = text.replace("\r", "\n")
    text = re.sub(r"[\x01-\x08\x0B\x0C\x0E-\x1F\x7F]", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"[ \t]{2,}", " ", text)
    return text.strip()


def split_text(text: str) -> list[str]:
    paragraphs = [paragraph.strip() for paragraph in text.split("\n\n") if paragraph.strip()]
    chunks: list[str] = []
    current = ""

    for paragraph in paragraphs:
        candidate = f"{current}\n\n{paragraph}".strip() if current else paragraph
        if len(candidate) <= MAX_CHUNK_CHARS:
            current = candidate
            continue

        if current:
            chunks.append(current)
            current = ""

        if len(paragraph) <= MAX_CHUNK_CHARS:
            current = paragraph
            continue

        chunks.extend(split_long_paragraph(paragraph))

    if current:
        chunks.append(current)

    return chunks


def split_long_paragraph(paragraph: str) -> list[str]:
    sentences = re.split(r"(?<=[.!?。！？])\s+", paragraph)
    chunks: list[str] = []
    current = ""

    for sentence in sentences:
        sentence = sentence.strip()
        if not sentence:
            continue

        candidate = f"{current} {sentence}".strip() if current else sentence
        if len(candidate) <= MAX_CHUNK_CHARS:
            current = candidate
            continue

        if current:
            chunks.append(current)

        if len(sentence) <= MAX_CHUNK_CHARS:
            current = sentence
        else:
            chunks.extend(
                sentence[index:index + MAX_CHUNK_CHARS].strip()
                for index in range(0, len(sentence), MAX_CHUNK_CHARS)
                if sentence[index:index + MAX_CHUNK_CHARS].strip()
            )
            current = ""

    if current:
        chunks.append(current)

    return chunks


def estimate_token_count(text: str) -> int:
    return max(1, math.ceil(len(text.split()) * 1.3))


def extract_section_title(text: str) -> str | None:
    first_line = text.splitlines()[0].strip()
    return first_line[:120] if first_line else None


def build_preview_text(text: str) -> str:
    try:
        summary = summarize_document(text)
        if summary:
            return summary
    except Exception:
        pass

    return text[:300]


def build_section_title(document_title: str, chunk_no: int, text: str) -> str | None:
    extracted_title = extract_section_title(text)
    if extracted_title and is_readable_title(extracted_title):
        return extracted_title[:120]

    return f"{document_title} - Chunk {chunk_no}"


def is_readable_title(text: str) -> bool:
    candidate = text.strip()
    if not candidate:
        return False

    if len(candidate) < 2:
        return False

    allowed = sum(
        1 for ch in candidate
        if ch.isalnum() or ch.isspace() or ch in "-_:/()[]{}.,"
    )
    ratio = allowed / len(candidate)

    if ratio < 0.6:
        return False

    weird_markers = ["�", "\u0000", "硫", "??", "Ã", "ð", "ì", "ë", "ê", "í", "ã", "â"]
    if any(marker in candidate for marker in weird_markers):
        return False

    return True


def repair_mojibake_text(text: str) -> str:
    candidate = text.strip()
    if not candidate:
        return text

    suspicious_chars = sum(1 for ch in candidate if 0x00C0 <= ord(ch) <= 0x00FF)
    ratio = suspicious_chars / max(1, len(candidate))

    if ratio < 0.05:
        return text

    try:
        repaired = candidate.encode("latin1", errors="ignore").decode("utf-8", errors="ignore")
    except Exception:
        return text

    if not repaired.strip():
        return text

    repaired_suspicious = sum(1 for ch in repaired if 0x00C0 <= ord(ch) <= 0x00FF)
    if repaired_suspicious < suspicious_chars:
        return repaired

    return text


def is_low_quality_extraction(text: str) -> bool:
    candidate = (text or "").strip()
    if not candidate:
        return True

    sample = candidate[:1000]
    weird_chars = sum(
        1 for ch in sample
        if not (ch.isalnum() or ch.isspace() or ch in ".,!?;:'\"-_/()[]{}<>@#%&*+=~|")
    )
    weird_ratio = weird_chars / max(1, len(sample))

    gibberish_markers = ["�", "硫", "Ã", "ð", "¿", "½", "¼", " IZ5 ", " d'N^ "]
    if any(marker in sample for marker in gibberish_markers):
        return True

    return weird_ratio > 0.35


def make_hash_embedding(text: str, dimension: int) -> list[float]:
    vector = [0.0] * dimension
    tokens = re.findall(r"[\w가-힣]+", text.lower())

    if not tokens:
        return vector

    for token in tokens:
        digest = hashlib.sha256(token.encode("utf-8")).digest()
        index = int.from_bytes(digest[:4], "big") % dimension
        sign = 1.0 if digest[4] % 2 == 0 else -1.0
        weight = 1.0 + (digest[5] / 255.0)
        vector[index] += sign * weight

    norm = math.sqrt(sum(value * value for value in vector))
    if norm == 0:
        return vector

    return [round(value / norm, 6) for value in vector]
